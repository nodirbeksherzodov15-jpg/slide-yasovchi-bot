"""PowerPoint (.pptx) faylini yig'ish. Telegramga bog'liq emas — alohida sinash mumkin."""
import copy
from concurrent.futures import ThreadPoolExecutor

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from images import fetch_image

FONT = "Calibri"
MAX_IMAGES = 5
IMG_W, IMG_H = 4.9, 3.9          # dyuym
NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

THEMES = {
    "1": ("Universitet", "123B63", "2E75B6", "F3F7FB", "FFFFFF"),
    "2": ("Qora-oltin", "20212A", "D6B15D", "F4F0E6", "FFFFFF"),
    "3": ("Gradient ranglar", "4934A4", "E65B91", "F8F5FF", "FFFFFF"),
    "4": ("Data va statistika", "12343B", "0F9D8A", "EEF7F5", "FFFFFF"),
    "5": ("Minimalistik", "343A40", "7C8A96", "F7F7F5", "FFFFFF"),
    "6": ("Texnologik ko‘k", "071D49", "00A6FB", "EDF6FF", "FFFFFF"),
    "7": ("Iqtisodiyot", "174A35", "6CA66A", "F1F7F0", "FFFFFF"),
    "8": ("Tarix", "5B3825", "B47B48", "F8F0E4", "FFFFFF"),
    "9": ("Tibbiyot", "174B68", "37A6A0", "EFF9FA", "FFFFFF"),
    "10": ("Tabiat", "285943", "8DBF67", "F1F7E9", "FFFFFF"),
    "11": ("Kreativ", "3B2063", "F08A4B", "FFF5EF", "FFFFFF"),
    "12": ("Klassik akademik", "27364B", "9A8058", "F5F2EA", "FFFFFF"),
}


def rgb(h):
    return RGBColor.from_string(h)


def _luma(h):
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return 0.299 * r + 0.587 * g + 0.114 * b


def add_text(slide, text, x, y, w, h, size=22, color="263238", bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = rgb(color)
    return box


def add_bullets(slide, bullets, x, y, w, h, theme, size=None):
    _, dark, accent, _, _ = theme
    total = sum(len(b) for b in bullets)
    # kenglikka qarab shrift o'lchamini tanlaymiz (matn qutidan chiqib ketmasin)
    if size is None:
        area = w * h
        size = 24
        for s in (24, 22, 20, 18, 16, 14):
            chars_per_line = max(10, int(w * 72 / (s * 0.52)))
            lines = sum(-(-len(b) // chars_per_line) for b in bullets)
            need = lines * s * 1.25 / 72 + len(bullets) * 0.16
            size = s
            if need <= h:
                break
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.55)
        m = p.add_run()
        m.text = "■  "
        m.font.name, m.font.size = FONT, Pt(size * 0.6)
        m.font.color.rgb = rgb(accent)
        t = p.add_run()
        t.text = b
        t.font.name, t.font.size = FONT, Pt(size)
        t.font.color.rgb = rgb(dark)
    return box


def add_base(prs, theme):
    _, dark, accent, light, _ = theme
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(light)
    for (x, y, w, h) in [(0, 0, 0.20, prs.slide_height / 914400), (0.55, 0.48, 12.0, 0.05)]:
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(accent)
        s.line.fill.background()
    return slide


def add_footer(slide, num, theme, label):
    _, dark, accent, _, _ = theme
    add_text(slide, label, 0.65, 7.08, 9.0, 0.28, 10, dark)
    add_text(slide, str(num), 12.0, 7.06, 0.5, 0.3, 11, accent, True, PP_ALIGN.RIGHT)


def add_fade(slide):
    """Haqiqiy 'fade' slayd o'tish effekti (OOXML)."""
    sld = slide._element
    for old in sld.findall(f"{{{NS}}}transition"):
        sld.remove(old)
    tr = etree.SubElement(sld, f"{{{NS}}}transition")
    tr.set("spd", "med")
    etree.SubElement(tr, f"{{{NS}}}fade")
    # tartib: cSld, clrMapOvr, transition, timing
    clr = sld.find(f"{{{NS}}}clrMapOvr")
    if clr is not None:
        sld.remove(tr)
        clr.addnext(tr)


def set_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _title_size(text, base=34):
    n = len(text)
    return base if n <= 40 else base - 4 if n <= 70 else base - 8 if n <= 110 else base - 12


def create_presentation(data, content, output_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    theme = THEMES[data["theme"]]
    theme_name, dark, accent, light, white = theme
    topic, subject = data["topic"], data["subject"]
    author = f'{data["name"]} {data["surname"]}'.strip()
    label = f"{topic[:60]} • {author}"
    slides_data = content["slides"]

    # --- rasmlarni parallel yuklab olamiz (eng ko'pi bilan MAX_IMAGES ta)
    wanted = [(i, s["image_query"]) for i, s in enumerate(slides_data) if s.get("image_query")][:MAX_IMAGES]
    images = {}
    if wanted:
        with ThreadPoolExecutor(max_workers=3) as ex:
            futs = {i: ex.submit(fetch_image, q, IMG_W / IMG_H) for i, q in wanted}
            for i, fu in futs.items():
                try:
                    res = fu.result(timeout=60)
                except Exception:
                    res = None
                if res:
                    images[i] = res

    # --- Muqova
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(dark)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.72), prs.slide_width, Inches(0.78))
    band.fill.solid()
    band.fill.fore_color.rgb = rgb(accent)
    band.line.fill.background()
    # accent rangi qorong'i fonda o'qilishi uchun
    sub_color = accent if _luma(accent) > 90 else "D9E2F2"
    add_text(s, subject.upper(), 0.85, 0.8, 11.5, 0.6, 18, sub_color, True)
    add_text(s, topic, 0.85, 1.6, 11.6, 2.6, _title_size(topic), white, True)
    add_text(s, f"Tayyorladi: {author}", 0.9, 4.6, 10.8, 0.55, 20, white)
    add_text(s, f"Dizayn: {theme_name}", 0.9, 5.25, 10.8, 0.4, 12, "D9E2F2")
    add_fade(s)

    # --- Reja
    s = add_base(prs, theme)
    add_text(s, "REJA", 0.75, 0.72, 11.7, 0.65, 30, dark, True)
    plan = content["plan"][:5] or ["Kirish", "Asosiy qism", "Xulosa"]
    step = min(1.15, 4.9 / len(plan))
    for i, item in enumerate(plan):
        y = 1.75 + i * step
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(y), Inches(11.2), Inches(step * 0.7))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(white)
        sh.line.color.rgb = rgb(accent)
        add_text(s, f"{i+1:02d}", 1.15, y + step * 0.08, 0.7, 0.45, 20, accent, True)
        add_text(s, item, 2.0, y + step * 0.08, 9.8, 0.5, 20, dark, True)
    add_footer(s, 2, theme, label)
    add_fade(s)

    # --- Asosiy slaydlar
    for i, sd in enumerate(slides_data):
        num = i + 3
        s = add_base(prs, theme)
        add_text(s, sd["title"], 0.75, 0.68, 11.8, 0.8, _title_size(sd["title"], 28), dark, True)
        if i in images:
            img, credit = images[i]
            s.shapes.add_picture(img, Inches(7.6), Inches(1.75), width=Inches(IMG_W), height=Inches(IMG_H))
            add_text(s, credit, 7.6, 1.75 + IMG_H + 0.05, IMG_W, 0.5, 9, "5F6B73")
            add_bullets(s, sd["bullets"], 0.85, 1.75, 6.5, 4.9, theme)
        else:
            add_bullets(s, sd["bullets"], 0.85, 1.75, 11.6, 4.9, theme)
        add_footer(s, num, theme, label)
        set_notes(s, sd.get("notes"))
        add_fade(s)

    # --- Xulosa
    s = add_base(prs, theme)
    add_text(s, "XULOSA", 0.75, 0.72, 11.7, 0.7, 30, dark, True)
    add_bullets(s, content["conclusion"] or ["Mavzu bo'yicha asosiy fikrlar umumlashtirildi."],
                0.95, 1.8, 11.2, 3.6, theme)
    add_text(s, f"Tayyorladi: {author}", 0.95, 5.9, 10.0, 0.5, 18, accent if _luma(accent) < 190 else dark, True)
    add_footer(s, len(prs.slides), theme, label)
    add_fade(s)

    prs.save(output_path)
    return len(prs.slides), len(images)
